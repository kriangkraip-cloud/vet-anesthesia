"""Pharmacy drug stock: items, receive/dispense ledger, usage from anesthetic records,
and Excel exports (monthly Category-2 narcotic report + drug usage log).

Stock on hand for an item =
    opening + received + positive adjustments
  - manually dispensed - negative adjustments
  - quantities taken from drug entries in anesthetic records (matched by `link_drug_name`)

Record usage is *computed on the fly* from the drug entries, never copied, so editing or
deleting a drug entry in a record is always reflected in the stock figures.
"""
import io
import re
from calendar import monthrange
from collections import OrderedDict
from datetime import date, datetime, timedelta
from typing import Dict, List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.properties import PageSetupProperties
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from .. import auth, models, schemas
from ..database import get_db

router = APIRouter(prefix="/api/stock", tags=["stock"])

_TH_OFFSET = timedelta(hours=7)
USAGE_BASES = ("volume_ml", "dose_mg", "dose_mcg", "per_entry")
TX_TYPES = ("opening", "receive", "dispense", "adjust")
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
THAI_MONTHS = ["มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
               "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"]

# Fixed licence / premises text of the official report. Per the hospital's form, everything
# here stays the same every month — only the month and the Buddhist-era year change.
REPORT_TITLE = "รายงานรับ จ่ายยาเสพติดให้โทษประเภท ๒"
REPORT_LICENCE_NOTE = "ใบอนุญาตเกี่ยวกับยาเสพติดให้โทษในประเภท ๒ (ให้เลือกข้อความใน        เพียงช่องเดียว)"
REPORT_LICENCE_CHOICE = "☑ จำหน่าย                     ☐ มีไว้ครอบครอง"
REPORT_LICENSEE = ("ชื่อผู้รับอนุญาต..............นายเชาวพันธ์.......ยินหาญมิ่งมงคล......................................... "
                   "ใบอนุญาตเลขที่...........จยส..2-ร 377.../..2565................................................................")
REPORT_PREMISES = "สถานที่ชื่อ...โรงพยาบาลสัตว์แอนิมอลสเปซ" + "." * 150
REPORT_ADDRESS_1 = ("ตั้งอยู่เลขที่......99/30.........หมู่ที่.................. ตรอก/ซอย.................................... "
                    "ถนน........พุทธมณฑลสาย...2....ซอย..24............................ตำบล/แขวง............ศาลาธรรมสพน์.........................")
REPORT_ADDRESS_2 = ("อำเภอ/เขต..........ทวีวัฒนา................ จังหวัด......กรุงเทพมหานคร.........รหัสไปรษณีย์......10160...... "
                    "โทรศัพท์....021162214.......... โทรสาร............................ E-mail..................................")
REPORT_INTRO = "ขอรายงานผลการดำเนินกิจการเกี่ยวกับยาเสพติดให้โทษในประเภท ๒ ดังนี้"
REPORT_FOOTNOTES = [
    '          หมายเหตุ : (๑) * ระบุหน่วย เช่น กรณียาน้ำให้ระบุเป็น "มิลลิลิตร" หรือ กรณียาเม็ดให้ระบุเป็น "เม็ด" '
    'หรือ "แคปซูล" หรือ กรณียาฉีดให้ระบุเป็น "ampule" หรือ "vial" ฯลฯ',
    "                             (๒) ** โปรดลงชื่อ",
    "                             (๓) ให้ขีดฆ่าข้อความที่ไม่ต้องการออก",
]


# ── Small helpers ────────────────────────────────────────────────────────────

def _today_th() -> date:
    return (datetime.utcnow() + _TH_OFFSET).date()


def _clean(v: Optional[str]) -> Optional[str]:
    v = (v or "").strip()
    return v or None


def _num(v: Optional[float]) -> Optional[float]:
    """Round away float noise (0.1+0.2) so quantities print cleanly."""
    return None if v is None else round(v, 4)


def _parse_period(month: Optional[str], date_from: Optional[date], date_to: Optional[date]):
    """Return (d_from, d_to, month_mode). `month` (YYYY-MM) wins over an explicit range."""
    if month:
        m = re.fullmatch(r"(\d{4})-(\d{2})", month)
        if not m or not 1 <= int(m.group(2)) <= 12:
            raise HTTPException(status_code=400, detail="month must be YYYY-MM")
        y, mo = int(m.group(1)), int(m.group(2))
        return date(y, mo, 1), date(y, mo, monthrange(y, mo)[1]), True
    if not date_from or not date_to:
        raise HTTPException(status_code=400, detail="Provide month=YYYY-MM or both date_from and date_to")
    if date_from > date_to:
        raise HTTPException(status_code=400, detail="date_from must not be after date_to")
    full_month = (date_from.day == 1 and date_to == date(date_to.year, date_to.month, monthrange(date_to.year, date_to.month)[1])
                  and (date_from.year, date_from.month) == (date_to.year, date_to.month))
    return date_from, date_to, full_month


def _require_item(db: Session, item_id: int) -> models.DrugStockItem:
    item = db.query(models.DrugStockItem).filter(models.DrugStockItem.id == item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Stock item not found")
    return item


def _check_basis(basis: Optional[str]):
    if basis is not None and basis not in USAGE_BASES:
        raise HTTPException(status_code=400, detail=f"usage_basis must be one of {', '.join(USAGE_BASES)}")


def _display_patient(p: Optional[models.Patient]) -> str:
    return f"{p.name} / {p.hn}" if p else ""


# ── Quantity of a drug entry ─────────────────────────────────────────────────

def _entry_quantities(entry: models.DrugEntry, weight: Optional[float]) -> Dict[str, Optional[float]]:
    """Best-effort volume (mL) and total dose (mg) of one drug entry."""
    dose, unit = entry.dose, (entry.dose_unit or "")
    conc = entry.concentration
    if conc is not None and (entry.concentration_unit or "") == "mcg/mL":
        conc = conc / 1000.0

    mg = None
    if dose is not None:
        if unit in ("total", "mg"):
            mg = dose
        elif unit == "mg/kg" and weight:
            mg = dose * weight
        elif unit == "mcg/kg" and weight:
            mg = dose * weight / 1000.0
        elif unit == "mL" and conc:
            mg = dose * conc
    if mg is None and entry.calculated_volume is not None and conc:
        mg = entry.calculated_volume * conc
    if mg is None and entry.calculated_total_dose is not None:
        mg = entry.calculated_total_dose

    ml = None
    if entry.calculated_volume is not None:
        ml = entry.calculated_volume
    elif dose is not None and unit == "mL":
        ml = dose
    elif mg is not None and conc:
        ml = mg / conc
    return {"volume_ml": ml, "dose_mg": mg}


def _stock_quantity(item: models.DrugStockItem, entry: models.DrugEntry, weight: Optional[float]) -> Optional[float]:
    """Quantity in the item's own stock unit, or None when the entry has no usable amount."""
    basis = item.usage_basis or "volume_ml"
    if basis == "per_entry":
        raw = 1.0
    else:
        q = _entry_quantities(entry, weight)
        if basis == "dose_mcg":
            raw = q["dose_mg"] * 1000.0 if q["dose_mg"] is not None else None
        else:
            raw = q.get(basis)
    if raw is None:
        return None
    divisor = item.usage_divisor if item.usage_divisor and item.usage_divisor > 0 else 1.0
    return _num(raw / divisor)


# ── Usage coming from anesthetic records ─────────────────────────────────────

def _record_usage(db: Session, item: models.DrugStockItem, d_from: Optional[date], d_to: Optional[date]) -> List[dict]:
    """Drug entries in anesthetic records that belong to this stock item, within [d_from, d_to]."""
    link = (item.link_drug_name or "").strip().lower()
    if not link:
        return []
    lower_bound = d_from
    if item.track_start and (lower_bound is None or item.track_start > lower_bound):
        lower_bound = item.track_start

    q = (
        db.query(models.DrugEntry, models.AnestheticRecord, models.Patient)
        .join(models.AnestheticRecord, models.DrugEntry.record_id == models.AnestheticRecord.id)
        .join(models.Patient, models.AnestheticRecord.patient_id == models.Patient.id)
        .filter(func.lower(func.trim(models.DrugEntry.drug_name)) == link)
        .filter(models.AnestheticRecord.record_date.isnot(None))
    )
    if lower_bound:
        q = q.filter(models.AnestheticRecord.record_date >= lower_bound)
    if d_to:
        q = q.filter(models.AnestheticRecord.record_date <= d_to)

    rows = []
    for entry, rec, pat in q.order_by(models.AnestheticRecord.record_date, models.DrugEntry.id).all():
        qty = _stock_quantity(item, entry, rec.weight_at_record or pat.weight)
        rows.append({
            "date": rec.record_date,
            "record_id": rec.id,
            "entry_id": entry.id,
            "patient_id": pat.id,
            "hn": pat.hn,
            "dispensed_to": _display_patient(pat),
            "quantity": qty,
        })
    return rows


# ── Ledger ───────────────────────────────────────────────────────────────────

_ORDER = {"opening": 0, "receive": 1, "adjust": 2, "dispense": 3, "record": 4}


def _tx_signed(tx: models.DrugStockTransaction) -> float:
    return -tx.quantity if tx.tx_type == "dispense" else tx.quantity


def _balance_as_of(db: Session, item: models.DrugStockItem, d: date) -> float:
    """Stock on hand at the *end* of day `d`."""
    manual = sum(
        _tx_signed(t) for t in db.query(models.DrugStockTransaction)
        .filter(models.DrugStockTransaction.item_id == item.id, models.DrugStockTransaction.tx_date <= d).all()
    )
    used = sum(r["quantity"] or 0 for r in _record_usage(db, item, None, d))
    return _num(manual - used)


def _build_ledger(db: Session, item: models.DrugStockItem, d_from: date, d_to: date) -> dict:
    txs = (db.query(models.DrugStockTransaction)
           .filter(models.DrugStockTransaction.item_id == item.id)
           .order_by(models.DrugStockTransaction.tx_date, models.DrugStockTransaction.id).all())
    usage = _record_usage(db, item, None, d_to)

    before = [t for t in txs if t.tx_date < d_from]
    carried = _num(sum(_tx_signed(t) for t in before) - sum(r["quantity"] or 0 for r in usage if r["date"] < d_from))
    # "tracking" = the system holds real stock data (an opening balance or a receipt) up to the end of the period
    tracking = any(t.tx_type in ("opening", "receive") and t.tx_date <= d_to for t in txs)

    cur_batch, cur_mfr = None, item.manufacturer
    for t in before:
        if t.tx_type in ("opening", "receive"):
            cur_batch = t.batch_no or cur_batch
            cur_mfr = t.manufacturer or cur_mfr

    start_batch, start_mfr = cur_batch, cur_mfr

    raw_rows = []
    for t in txs:
        if d_from <= t.tx_date <= d_to:
            raw_rows.append((t.tx_date, _ORDER[t.tx_type], t.id, "tx", t))
    for r in usage:
        if d_from <= r["date"] <= d_to:
            raw_rows.append((r["date"], _ORDER["record"], r["entry_id"], "record", r))
    raw_rows.sort(key=lambda x: (x[0], x[1], x[2]))

    rows, balance = [], carried
    total_opening = total_in = total_out = 0.0
    for d, _, _, kind, obj in raw_rows:
        row = {"date": d, "kind": None, "tx_id": None, "record_id": None, "entry_id": None,
               "source": None, "dispensed_to": None, "batch_no": None, "manufacturer": None,
               "qty_opening": 0.0, "qty_in": 0.0, "qty_out": 0.0, "qty_unknown": False, "note": None}
        if kind == "tx":
            t = obj
            row.update(kind=t.tx_type, tx_id=t.id, note=t.note, dispensed_to=t.dispensed_to,
                       source=t.source, patient_id=t.patient_id)
            if t.tx_type == "opening":
                row["qty_opening"] = t.quantity
                total_opening += t.quantity
            elif t.tx_type == "receive" or (t.tx_type == "adjust" and t.quantity >= 0):
                row["qty_in"] = t.quantity
                total_in += t.quantity
            else:  # dispense or negative adjust
                row["qty_out"] = abs(t.quantity)
                total_out += abs(t.quantity)
            if t.tx_type in ("opening", "receive"):
                cur_batch = t.batch_no or cur_batch
                cur_mfr = t.manufacturer or cur_mfr
                row["batch_no"] = t.batch_no
                row["manufacturer"] = t.manufacturer or item.manufacturer
            else:
                row["batch_no"] = t.batch_no or cur_batch
                row["manufacturer"] = t.manufacturer or cur_mfr
            if t.tx_type == "adjust":
                row["source"] = row["source"] or "ปรับปรุงยอด"
        else:
            r = obj
            row.update(kind="record", record_id=r["record_id"], entry_id=r["entry_id"],
                       dispensed_to=r["dispensed_to"], patient_id=r["patient_id"],
                       batch_no=cur_batch, manufacturer=cur_mfr)
            if r["quantity"] is None:
                row["qty_unknown"] = True
            else:
                row["qty_out"] = r["quantity"]
                total_out += r["quantity"]
        balance = _num(balance + row["qty_opening"] + row["qty_in"] - row["qty_out"])
        row["balance"] = balance
        rows.append(row)

    return {
        "item_id": item.id, "item_name": item.name, "unit": item.unit,
        "date_from": d_from, "date_to": d_to,
        "carried": carried, "tracking": tracking, "rows": rows,
        "start_batch": start_batch, "start_mfr": start_mfr,
        "total_opening": _num(total_opening), "total_in": _num(total_in), "total_out": _num(total_out),
        "closing": balance,
    }


def _ledger_json(led: dict) -> dict:
    out = dict(led)
    out["date_from"], out["date_to"] = led["date_from"].isoformat(), led["date_to"].isoformat()
    out["rows"] = [{**r, "date": r["date"].isoformat()} for r in led["rows"]]
    return out


def _item_json(db: Session, item: models.DrugStockItem, today: Optional[date] = None) -> dict:
    today = today or _today_th()
    m_from, m_to = today.replace(day=1), today
    led = _build_ledger(db, item, m_from, m_to)
    return {
        "id": item.id, "name": item.name, "unit": item.unit, "is_controlled": bool(item.is_controlled),
        "manufacturer": item.manufacturer, "link_drug_name": item.link_drug_name,
        "usage_basis": item.usage_basis, "usage_divisor": item.usage_divisor,
        "track_start": item.track_start.isoformat() if item.track_start else None,
        "is_active": bool(item.is_active), "notes": item.notes,
        "balance": _balance_as_of(db, item, today),
        "month_in": _num(led["total_opening"] + led["total_in"]),
        "month_out": led["total_out"],
        "unknown_usage": sum(1 for r in led["rows"] if r["qty_unknown"]),
    }


# ── Items ────────────────────────────────────────────────────────────────────

@router.get("/items")
async def list_items(
    q: Optional[str] = Query(None),
    include_inactive: bool = Query(False),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_pharmacy_or_admin),
):
    query = db.query(models.DrugStockItem)
    if not include_inactive:
        query = query.filter(models.DrugStockItem.is_active.is_(True))
    if q:
        like = f"%{q.strip()}%"
        query = query.filter(or_(models.DrugStockItem.name.ilike(like),
                                 models.DrugStockItem.link_drug_name.ilike(like),
                                 models.DrugStockItem.manufacturer.ilike(like)))
    today = _today_th()
    return [_item_json(db, i, today) for i in query.order_by(models.DrugStockItem.name).all()]


@router.post("/items")
async def create_item(
    data: schemas.StockItemCreate,
    db: Session = Depends(get_db),
    user: models.User = Depends(auth.get_current_pharmacy_or_admin),
):
    name = _clean(data.name)
    if not name:
        raise HTTPException(status_code=400, detail="Drug name is required")
    _check_basis(data.usage_basis)
    if data.usage_divisor is not None and data.usage_divisor <= 0:
        raise HTTPException(status_code=400, detail="usage_divisor must be greater than 0")
    if db.query(models.DrugStockItem).filter(func.lower(models.DrugStockItem.name) == name.lower()).first():
        raise HTTPException(status_code=400, detail="A stock item with this name already exists")

    if data.opening_qty is not None and data.opening_qty < 0:
        raise HTTPException(status_code=400, detail="Opening quantity cannot be negative")
    track_start = data.track_start or _today_th()
    item = models.DrugStockItem(
        name=name, unit=_clean(data.unit) or "mL", is_controlled=data.is_controlled,
        manufacturer=_clean(data.manufacturer), link_drug_name=_clean(data.link_drug_name),
        usage_basis=data.usage_basis, usage_divisor=data.usage_divisor or 1.0,
        track_start=track_start, notes=_clean(data.notes), is_active=True,
    )
    db.add(item)
    db.flush()
    if data.opening_qty:
        db.add(models.DrugStockTransaction(
            item_id=item.id, tx_date=track_start, tx_type="opening", quantity=data.opening_qty,
            batch_no=_clean(data.opening_batch), manufacturer=item.manufacturer,
            source="ยอดยกมา", created_by_id=user.id,
        ))
    db.commit()
    db.refresh(item)
    return _item_json(db, item)


@router.put("/items/{item_id}")
async def update_item(
    item_id: int,
    data: schemas.StockItemUpdate,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_pharmacy_or_admin),
):
    item = _require_item(db, item_id)
    changes = data.model_dump(exclude_unset=True)
    if "usage_basis" in changes:
        _check_basis(changes["usage_basis"])
    if changes.get("usage_divisor") is not None and changes["usage_divisor"] <= 0:
        raise HTTPException(status_code=400, detail="usage_divisor must be greater than 0")
    if "name" in changes:
        name = _clean(changes["name"])
        if not name:
            raise HTTPException(status_code=400, detail="Drug name is required")
        dup = db.query(models.DrugStockItem).filter(
            func.lower(models.DrugStockItem.name) == name.lower(), models.DrugStockItem.id != item_id).first()
        if dup:
            raise HTTPException(status_code=400, detail="A stock item with this name already exists")
        changes["name"] = name
    for field in ("unit", "manufacturer", "link_drug_name", "notes"):
        if field in changes:
            changes[field] = _clean(changes[field])
    for field, value in changes.items():
        setattr(item, field, value)
    item.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(item)
    return _item_json(db, item)


@router.delete("/items/{item_id}")
async def delete_item(
    item_id: int,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_pharmacy_or_admin),
):
    item = _require_item(db, item_id)
    if db.query(models.DrugStockTransaction).filter(models.DrugStockTransaction.item_id == item_id).count():
        raise HTTPException(status_code=400,
                            detail="This item has stock movements. Deactivate it instead of deleting it.")
    db.delete(item)
    db.commit()
    return {"ok": True}


@router.get("/items/{item_id}/ledger")
async def item_ledger(
    item_id: int,
    month: Optional[str] = Query(None),
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_pharmacy_or_admin),
):
    item = _require_item(db, item_id)
    d_from, d_to, _month = _parse_period(month, date_from, date_to)
    return _ledger_json(_build_ledger(db, item, d_from, d_to))


# ── Transactions ─────────────────────────────────────────────────────────────

def _validate_tx(tx_type: str, quantity: float):
    if tx_type not in TX_TYPES:
        raise HTTPException(status_code=400, detail=f"tx_type must be one of {', '.join(TX_TYPES)}")
    if tx_type == "adjust":
        if not quantity:
            raise HTTPException(status_code=400, detail="Adjustment quantity cannot be 0")
    elif quantity is None or quantity <= 0:
        raise HTTPException(status_code=400, detail="Quantity must be greater than 0")


def _tx_json(t: models.DrugStockTransaction) -> dict:
    return {
        "id": t.id, "item_id": t.item_id, "tx_date": t.tx_date.isoformat(), "tx_type": t.tx_type,
        "quantity": t.quantity, "batch_no": t.batch_no, "manufacturer": t.manufacturer, "source": t.source,
        "patient_id": t.patient_id, "dispensed_to": t.dispensed_to, "note": t.note,
    }


def _fill_dispensed_to(db: Session, patient_id: Optional[int], dispensed_to: Optional[str]) -> Optional[str]:
    dispensed_to = _clean(dispensed_to)
    if patient_id and not dispensed_to:
        pat = db.query(models.Patient).filter(models.Patient.id == patient_id).first()
        if not pat:
            raise HTTPException(status_code=404, detail="Patient not found")
        return _display_patient(pat)
    return dispensed_to


@router.post("/transactions")
async def create_transaction(
    data: schemas.StockTxCreate,
    db: Session = Depends(get_db),
    user: models.User = Depends(auth.get_current_pharmacy_or_admin),
):
    item = _require_item(db, data.item_id)
    _validate_tx(data.tx_type, data.quantity)
    tx = models.DrugStockTransaction(
        item_id=item.id, tx_date=data.tx_date or _today_th(), tx_type=data.tx_type, quantity=data.quantity,
        batch_no=_clean(data.batch_no), manufacturer=_clean(data.manufacturer), source=_clean(data.source),
        patient_id=data.patient_id, dispensed_to=_fill_dispensed_to(db, data.patient_id, data.dispensed_to),
        note=_clean(data.note), created_by_id=user.id,
    )
    db.add(tx)
    db.commit()
    db.refresh(tx)
    return _tx_json(tx)


@router.put("/transactions/{tx_id}")
async def update_transaction(
    tx_id: int,
    data: schemas.StockTxUpdate,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_pharmacy_or_admin),
):
    tx = db.query(models.DrugStockTransaction).filter(models.DrugStockTransaction.id == tx_id).first()
    if not tx:
        raise HTTPException(status_code=404, detail="Transaction not found")
    changes = data.model_dump(exclude_unset=True)
    _validate_tx(changes.get("tx_type", tx.tx_type), changes.get("quantity", tx.quantity))
    for field in ("batch_no", "manufacturer", "source", "dispensed_to", "note"):
        if field in changes:
            changes[field] = _clean(changes[field])
    for field, value in changes.items():
        if field == "tx_date" and value is None:
            continue
        setattr(tx, field, value)
    if "patient_id" in changes and not tx.dispensed_to:
        tx.dispensed_to = _fill_dispensed_to(db, tx.patient_id, None)
    db.commit()
    db.refresh(tx)
    return _tx_json(tx)


@router.delete("/transactions/{tx_id}")
async def delete_transaction(
    tx_id: int,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_pharmacy_or_admin),
):
    tx = db.query(models.DrugStockTransaction).filter(models.DrugStockTransaction.id == tx_id).first()
    if not tx:
        raise HTTPException(status_code=404, detail="Transaction not found")
    db.delete(tx)
    db.commit()
    return {"ok": True}


# ── Drug usage log (all drug entries in anesthetic records) ──────────────────

def _usage_log(db: Session, d_from: date, d_to: date, q: Optional[str] = None) -> dict:
    query = (
        db.query(models.DrugEntry, models.AnestheticRecord, models.Patient)
        .join(models.AnestheticRecord, models.DrugEntry.record_id == models.AnestheticRecord.id)
        .join(models.Patient, models.AnestheticRecord.patient_id == models.Patient.id)
        .filter(models.AnestheticRecord.record_date >= d_from, models.AnestheticRecord.record_date <= d_to)
        .filter(models.DrugEntry.drug_name.isnot(None))
    )
    if q:
        like = f"%{q.strip()}%"
        query = query.filter(or_(models.DrugEntry.drug_name.ilike(like), models.Patient.name.ilike(like),
                                 models.Patient.hn.ilike(like)))
    rows, summary = [], OrderedDict()
    for entry, rec, pat in query.order_by(models.AnestheticRecord.record_date, models.AnestheticRecord.id,
                                          models.DrugEntry.id).all():
        weight = rec.weight_at_record or pat.weight
        qty = _entry_quantities(entry, weight)
        when = (entry.time + _TH_OFFSET).strftime("%H:%M") if entry.time else None
        if entry.entry_type == "cri" and entry.rate is not None:
            dose_text = f"CRI {entry.rate:g} {entry.rate_unit or 'mL/hr'}"
        elif entry.dose is not None:
            dose_text = f"{entry.dose:g} {entry.dose_unit or ''}".strip()
        else:
            dose_text = ""
        rows.append({
            "date": rec.record_date.isoformat(), "time": when, "record_id": rec.id,
            "hn": pat.hn, "patient_name": pat.name, "species": pat.species, "weight": weight,
            "drug_name": entry.drug_name.strip(), "entry_type": entry.entry_type, "dose": dose_text,
            "volume_ml": _num(qty["volume_ml"]), "dose_mg": _num(qty["dose_mg"]), "route": entry.route,
            "anesthesiologist": rec.anesthesiologist, "surgeon": rec.surgeon,
        })
        key = entry.drug_name.strip().lower()
        s = summary.setdefault(key, {"drug_name": entry.drug_name.strip(), "count": 0, "volume_ml": 0.0, "dose_mg": 0.0})
        s["count"] += 1
        s["volume_ml"] += qty["volume_ml"] or 0
        s["dose_mg"] += qty["dose_mg"] or 0
    summ = sorted(({**s, "volume_ml": _num(s["volume_ml"]), "dose_mg": _num(s["dose_mg"])} for s in summary.values()),
                  key=lambda s: s["drug_name"].lower())
    return {"date_from": d_from.isoformat(), "date_to": d_to.isoformat(), "rows": rows, "summary": summ}


@router.get("/usage")
async def drug_usage(
    month: Optional[str] = Query(None),
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    q: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_pharmacy_or_admin),
):
    d_from, d_to, _month = _parse_period(month, date_from, date_to)
    return _usage_log(db, d_from, d_to, q)


# ── Excel exports ────────────────────────────────────────────────────────────

_THIN = Side(style="thin", color="000000")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_FONT = "TH Sarabun New"


def _font(size=14, bold=False):
    return Font(name=_FONT, size=size, bold=bold)


def _xlsx_response(wb: Workbook, filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type=XLSX_MIME,
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def _sheet_title(name: str, used: set) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "-", name).strip()[:28] or "Sheet"
    title, n = base, 2
    while title in used:
        title = f"{base[:26]}-{n}"
        n += 1
    used.add(title)
    return title


def _thai_date_text(d: date) -> str:
    return f"{d.day} {THAI_MONTHS[d.month - 1]} พ.ศ. {d.year + 543}"


def _write_narcotic_sheet(ws, item: models.DrugStockItem, led: dict, month_mode: bool):
    d_from, d_to = led["date_from"], led["date_to"]
    widths = {"A": 2, "B": 11, "C": 30, "D": 14, "E": 26, "F": 18, "G": 34,
              "H": 10, "I": 10, "J": 10, "K": 10, "L": 10, "M": 18}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    def line(row, text, size=14, bold=False, align="left", last_col="M"):
        ws.merge_cells(f"B{row}:{last_col}{row}")
        c = ws[f"B{row}"]
        c.value, c.font = text, _font(size, bold)
        c.alignment = Alignment(horizontal=align, vertical="center")

    line(2, REPORT_TITLE, 20, True, "center")
    if month_mode:
        line(3, f"ประจำเดือน.............{THAI_MONTHS[d_from.month - 1]}................. "
                f"พ.ศ. ..........{d_from.year + 543}..................", 16, True, "center")
    else:
        line(3, f"ระหว่างวันที่ {_thai_date_text(d_from)}  ถึงวันที่ {_thai_date_text(d_to)}", 16, True, "center")
    line(4, REPORT_LICENCE_NOTE)
    line(5, REPORT_LICENCE_CHOICE)
    line(6, REPORT_LICENSEE)
    line(7, REPORT_PREMISES)
    line(8, REPORT_ADDRESS_1)
    line(9, REPORT_ADDRESS_2)
    line(10, REPORT_INTRO)

    # table header (rows 11-12)
    for col, text in zip("BCDEFG", ["วัน  เดือน ปี", "ชื่อและความแรงของยา\nเสพติดให้โทษ\nในประเภท ๒",
                                    "เลขที่/\nรุ่นที่/ครั้งที่\nผลิต", "ชื่อผู้ผลิต/\nแหล่งผลิต", "ได้มาจาก", "จ่ายไปให้"]):
        ws.merge_cells(f"{col}11:{col}12")
        ws[f"{col}11"].value = text
    ws.merge_cells("H11:K11")
    ws["H11"].value = "จำนวน/ปริมาณยาเสพติดให้โทษในประเภท ๒"
    for col, text in zip("HIJK", ["ยอดยกมา", "รับ", "จ่าย", "คงเหลือ"]):
        ws[f"{col}12"].value = text
    for col, text in zip("LM", ["หน่วย *", "หมายเหตุ"]):
        ws.merge_cells(f"{col}11:{col}12")
        ws[f"{col}11"].value = text
    for row in (11, 12):
        for col in "BCDEFGHIJKLM":
            c = ws[f"{col}{row}"]
            c.font, c.border = _font(14, True), _BOX
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[11].height = 24
    ws.row_dimensions[12].height = 24

    # body
    tracking = led["tracking"]
    r = 13
    prev = {"name": None, "batch": None, "mfr": None}

    def ditto(col_key, value):
        if value and prev[col_key] == value:
            return '"'
        prev[col_key] = value
        return value or None

    def put(values, number_cols=()):
        nonlocal r
        for col, v in zip("BCDEFGHIJKLM", values):
            c = ws[f"{col}{r}"]
            c.value = v
            c.font, c.border = _font(14), _BOX
            if v == '"' or col in "BDL":
                horiz = "center"
            elif col in "HIJK":
                horiz = "right"
            else:
                horiz = "left"
            c.alignment = Alignment(horizontal=horiz, vertical="center", wrap_text=True)
        ws[f"B{r}"].number_format = "d/m/yy"
        r += 1

    first_date = d_from
    if tracking:   # carried-forward row; left blank when the system holds no stock data (per the form's rules)
        put([first_date, ditto("name", item.name), ditto("batch", led["start_batch"]),
             ditto("mfr", led["start_mfr"] or item.manufacturer),
             "ยอดยกมา", None, led["carried"], None, None, led["carried"], item.unit, None])
    for row in led["rows"]:
        if row["kind"] in ("opening",):
            src, h, i, j = "ยอดยกมา", row["qty_opening"], None, None
        elif row["kind"] == "receive":
            src, h, i, j = row["source"], None, row["qty_in"], None
        elif row["kind"] == "adjust":
            src, h, i, j = row["source"], None, row["qty_in"] or None, row["qty_out"] or None
        else:
            src, h, i, j = None, None, None, (row["qty_out"] if not row["qty_unknown"] else None)
        note = row["note"] if row["kind"] in ("adjust",) else None
        if row["qty_unknown"]:
            note = "ไม่ระบุปริมาณในใบบันทึก"
        elif row["kind"] == "dispense" and row["note"]:
            note = row["note"]
        put([row["date"], ditto("name", item.name), ditto("batch", row["batch_no"]),
             ditto("mfr", row["manufacturer"]), src,
             row["dispensed_to"] if row["kind"] in ("dispense", "record") else None,
             h, i, j, row["balance"] if tracking else None, item.unit, note])

    for _ in range(max(0, 5 - len(led["rows"]))):   # keep a few empty ruled rows like the paper form
        put([None] * 12)

    # totals
    tot_h = led["carried"] + led["total_opening"] if tracking else None
    ws[f"B{r}"].value = "รวม"
    for col, v in zip("HIJK", [_num(tot_h) if tot_h is not None else None, led["total_in"], led["total_out"],
                               led["closing"] if tracking else None]):
        ws[f"{col}{r}"].value = v
    for col in "BCDEFGHIJKLM":
        c = ws[f"{col}{r}"]
        c.font, c.border = _font(14, True), _BOX
        c.alignment = Alignment(horizontal="center" if col == "B" else "right", vertical="center")
    r += 1
    for text in REPORT_FOOTNOTES:
        line(r, text, 13)
        r += 1
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


@router.get("/export/narcotic")
async def export_narcotic(
    month: Optional[str] = Query(None),
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    item_ids: Optional[str] = Query(None, description="comma separated; default = all active controlled items"),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_pharmacy_or_admin),
):
    d_from, d_to, month_mode = _parse_period(month, date_from, date_to)
    if item_ids:
        try:
            ids = [int(x) for x in item_ids.split(",") if x.strip()]
        except ValueError:
            raise HTTPException(status_code=400, detail="item_ids must be a comma separated list of numbers")
        items = db.query(models.DrugStockItem).filter(models.DrugStockItem.id.in_(ids)).order_by(models.DrugStockItem.name).all()
    else:
        items = (db.query(models.DrugStockItem)
                 .filter(models.DrugStockItem.is_controlled.is_(True), models.DrugStockItem.is_active.is_(True))
                 .order_by(models.DrugStockItem.name).all())
    if not items:
        raise HTTPException(status_code=404, detail="No stock items to export. Mark items as Category-2 (controlled) or choose items.")

    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    for item in items:
        ws = wb.create_sheet(_sheet_title(item.name, used))
        _write_narcotic_sheet(ws, item, _build_ledger(db, item, d_from, d_to), month_mode)
    stamp = f"{d_from:%Y-%m}" if month_mode else f"{d_from:%Y%m%d}-{d_to:%Y%m%d}"
    return _xlsx_response(wb, f"narcotic_report_{stamp}.xlsx")


@router.get("/export/usage")
async def export_usage(
    month: Optional[str] = Query(None),
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    q: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_pharmacy_or_admin),
):
    d_from, d_to, month_mode = _parse_period(month, date_from, date_to)
    data = _usage_log(db, d_from, d_to, q)

    wb = Workbook()
    ws = wb.active
    ws.title = "Drug usage log"
    headers = ["Date", "Time", "HN", "Patient", "Species", "Weight (kg)", "Drug", "Type", "Dose", "Volume (mL)",
               "Total dose (mg)", "Route", "Anesthesiologist", "Surgeon", "Record #"]
    widths = [12, 8, 12, 22, 12, 11, 24, 18, 16, 12, 14, 10, 22, 22, 10]
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.border = Font(bold=True), _BOX
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    for n, row in enumerate(data["rows"], start=2):
        values = [date.fromisoformat(row["date"]), row["time"], row["hn"], row["patient_name"], row["species"],
                  row["weight"], row["drug_name"], row["entry_type"], row["dose"], row["volume_ml"], row["dose_mg"],
                  row["route"], row["anesthesiologist"], row["surgeon"], row["record_id"]]
        for i, v in enumerate(values, start=1):
            c = ws.cell(row=n, column=i, value=v)
            c.border = _BOX
        ws.cell(row=n, column=1).number_format = "dd/mm/yyyy"
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Summary by drug")
    for i, (h, w) in enumerate(zip(["Drug", "Times given", "Total volume (mL)", "Total dose (mg)"], [28, 12, 18, 16]), start=1):
        c = ws2.cell(row=1, column=i, value=h)
        c.font, c.border = Font(bold=True), _BOX
        ws2.column_dimensions[c.column_letter].width = w
    for n, s in enumerate(data["summary"], start=2):
        for i, v in enumerate([s["drug_name"], s["count"], s["volume_ml"], s["dose_mg"]], start=1):
            ws2.cell(row=n, column=i, value=v).border = _BOX
    stamp = f"{d_from:%Y-%m}" if month_mode else f"{d_from:%Y%m%d}-{d_to:%Y%m%d}"
    return _xlsx_response(wb, f"drug_usage_{stamp}.xlsx")
